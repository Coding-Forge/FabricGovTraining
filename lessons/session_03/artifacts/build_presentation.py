"""Build the Session 3 teaching deck from the six adjacent Markdown lessons.

Run: python lessons/session_03/artifacts/build_presentation.py
"""

from pathlib import Path
import re

from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "presentation" / "session_03_low_code_fabric_data_engineering.pptx"
SOURCES = sorted(ROOT.glob("0[1-6]_*.md"))
FONT_DIR = Path("/usr/share/fonts/truetype/dejavu")
FONT = "DejaVu Sans"
MONO = "DejaVu Sans Mono"
NAVY = "142D3B"
TEAL = "087F8C"
MINT = "DDF2EC"
INK = "183441"
GRAY = "526773"
PAPER = "F5F8FA"
WHITE = "FFFFFF"
GOLD = "B17622"
PALE_GOLD = "FFF1DA"
LINE = "DCE5EA"

prs = Presentation()
prs.slide_width = Inches(13.333333)
prs.slide_height = Inches(7.5)
prs.core_properties.title = "Session 3 | Low-code Fabric Data Engineering"
prs.core_properties.subject = "October 28, 2026 | SQL and Power Query practitioners"
prs.core_properties.author = "Brandon Campbell"
prs.core_properties.keywords = "Fabric, Warehouse, Dataflow Gen2, reconciliation, GCC High"
prs.core_properties.comments = "Based on the six Session 3 Markdown lessons."
module = "SESSION 3"
lesson = None
layout_records = []


def rgb(value):
    return RGBColor.from_string(value)


def rect(slide, x, y, w, h, fill, border=None):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.fill.background()
    if border:
        shape.line.color.rgb = rgb(border)
    return shape


def font_file(bold=False, mono=False):
    family = "DejaVuSansMono" if mono else "DejaVuSans"
    return FONT_DIR / f"{family}{'-Bold' if bold else ''}.ttf"


def wrap(value, width, size, bold=False, mono=False):
    font = ImageFont.truetype(str(font_file(bold, mono)), size * 4)
    max_width = width * 72 * 4
    result = []
    for line in str(value).split("\n"):
        current = ""
        for word in line.split(" "):
            candidate = f"{current} {word}" if current else word
            if font.getlength(candidate) > max_width and current:
                result.append(current)
                current = word
            else:
                current = candidate
        result.append(current)
    return result


def text(slide, value, x, y, w, h, size=22, color=INK, bold=False, mono=False):
    chosen = size
    while True:
        lines = wrap(value, w - 0.04, chosen, bold, mono)
        estimated = len(lines) * chosen * 1.22 / 72
        font = ImageFont.truetype(str(font_file(bold, mono)), chosen * 4)
        widest = max(font.getlength(line) for line in lines) / (72 * 4)
        if (estimated <= h - 0.04 and widest <= w - 0.04) or chosen <= 14:
            break
        chosen -= 1
    if estimated > h - 0.04 or widest > w - 0.04:
        raise ValueError(f"Text does not fit on slide {len(prs.slides)}: {value[:80]}")
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]
    p.text = "\v".join(lines)
    p.font.name = MONO if mono else FONT
    p.font.size = Pt(chosen)
    p.font.bold = bold
    p.font.color.rgb = rgb(color)
    p.line_spacing = 1.15
    p.space_after = Pt(0)
    layout_records.append((len(prs.slides), chosen, value[:70]))
    return shape


def notes(slide, value):
    source = f"\n\nSource: {lesson.name}" if lesson else ""
    slide.notes_slide.notes_text_frame.text = value + source


def base(title, subtitle=None, dark=False):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(NAVY if dark else PAPER)
    rect(slide, 0, 0, 13.333333, 0.09, TEAL)
    fg = WHITE if dark else INK
    text(slide, module.upper(), 0.65, 0.35, 12, 0.3, 11,
         MINT if dark else TEAL, True)
    text(slide, title, 0.65, 0.88, 12, 0.65, 30, fg, True)
    if subtitle:
        text(slide, subtitle, 0.65, 1.58, 12, 0.55, 17,
             "C8DCE3" if dark else GRAY)
    rect(slide, 0.65, 7.02, 12, 0.015, "355260" if dark else LINE)
    text(slide, "FABRIC TRAINING  /  OCTOBER 28, 2026",
         0.65, 7.13, 10.5, 0.2, 9, "C8DCE3" if dark else GRAY)
    text(slide, f"{len(prs.slides):02d}", 12.05, 7.09, 0.6, 0.28, 11, fg)
    return slide


def takeaway(slide, value):
    rect(slide, 0.65, 6.25, 12, 0.55, MINT)
    text(slide, value, 0.84, 6.35, 11.6, 0.32, 15, TEAL, True)


def cards(title, items, takeaway_text=None, speaker="", subtitle=None):
    slide = base(title, subtitle)
    count = len(items)
    width = (12 - 0.24 * (count - 1)) / count
    top = 2.35 if subtitle else 1.95
    height = 3.55 if subtitle else 3.95
    for i, (heading, body) in enumerate(items):
        x = 0.65 + i * (width + 0.24)
        rect(slide, x, top, width, height, WHITE)
        rect(slide, x, top, width, 0.055, TEAL)
        text(slide, heading, x + 0.21, top + 0.23, width - 0.42, 0.75, 22, TEAL, True)
        text(slide, body, x + 0.21, top + 1.12, width - 0.42, height - 1.32, 21)
    if takeaway_text:
        takeaway(slide, takeaway_text)
    notes(slide, speaker or title)
    return slide


def bullets(title, items, takeaway_text=None, speaker="", subtitle=None):
    slide = base(title, subtitle)
    top = 2.3 if subtitle else 1.95
    step = min(0.92, 4.0 / len(items))
    for i, item in enumerate(items):
        y = top + i * step
        text(slide, f"{i + 1:02d}", 0.68, y, 0.65, 0.5, 20, TEAL, True)
        text(slide, item, 1.48, y, 10.9, step - 0.08, 22)
    if takeaway_text:
        takeaway(slide, takeaway_text)
    notes(slide, speaker or title)
    return slide


def table(title, headers, rows, widths, speaker="", takeaway_text=None, size=18):
    slide = base(title)
    y = 1.95
    header_height = 0.58
    row_height = min(0.78, 3.9 / len(rows))
    for row_index, row in enumerate([headers] + rows):
        x = 0.65
        height = header_height if row_index == 0 else row_height
        for value, width in zip(row, widths):
            rect(slide, x, y, width - 0.025, height - 0.025,
                 NAVY if row_index == 0 else (WHITE if row_index % 2 else "EAF0F4"))
            text(slide, value, x + 0.12, y + 0.09, width - 0.25, height - 0.16,
                 size if row_index else 16, WHITE if row_index == 0 else INK,
                 row_index == 0)
            x += width
        y += height
    if takeaway_text:
        takeaway(slide, takeaway_text)
    notes(slide, speaker or title)
    return slide


def code(title, snippet, heading, explanation, speaker=""):
    slide = base(title)
    rect(slide, 0.65, 1.95, 7.55, 4.78, NAVY)
    text(slide, snippet, 0.9, 2.2, 7.05, 4.27, 18, WHITE, mono=True)
    text(slide, heading, 8.55, 2.05, 4.1, 0.9, 23, TEAL, True)
    text(slide, explanation, 8.55, 3.1, 4.1, 3.5, 21)
    notes(slide, speaker or title)
    return slide


def exercise(title, prompt, output, debrief, speaker=""):
    slide = base(title, "APPLY THE LESSON", dark=True)
    text(slide, prompt, 0.7, 2.45, 11.9, 1.5, 29, WHITE, True)
    rect(slide, 0.7, 4.45, 11.9, 1.65, "203F4D")
    text(slide, "PRODUCE", 0.95, 4.65, 1.5, 0.35, 13, MINT, True)
    text(slide, output, 2.55, 4.65, 9.7, 1.12, 21, WHITE)
    notes(slide, f"{speaker}\n\nDebrief / expected answer:\n{debrief}")
    return slide


def section(number, title, minutes, output):
    global module, lesson
    lesson = SOURCES[number - 1]
    module = f"{number:02d} / {title}"
    slide = base(title, f"LESSON {number}  /  {minutes} MINUTES", dark=True)
    text(slide, str(number).zfill(2), 0.7, 2.6, 3, 1.5, 86, MINT, True)
    text(slide, "WHAT YOU WILL PRODUCE", 4.25, 2.75, 8, 0.4, 13, MINT, True)
    text(slide, output, 4.25, 3.45, 8, 1.85, 31, WHITE, True)
    notes(slide, "Complete source lesson for presenter reference:\n\n" + lesson.read_text())


def build():
    global module, lesson
    slide = base("", dark=True)
    text(slide, "Low-code Fabric\nData Engineering",
         0.7, 0.95, 11.8, 1.5, 40, WHITE, True)
    text(slide, "SESSION 3", 0.7, 2.65, 11.8, 0.5, 19, MINT, True)
    text(slide, "From SharePoint extracts to a repeatable,\ngoverned reconciliation workload.",
         0.7, 3.25, 11.8, 1.4, 31, WHITE)
    text(slide, "October 28, 2026  |  11:00 AM-2:00 PM ET\nSQL and Power Query practitioners",
         0.7, 5.3, 11.8, 0.95, 20, "C8DCE3")
    notes(slide, "Open with the business outcome, not a product tour. This is the three-hour "
          "Session 3 workshop. All examples are synthetic. The deck covers all six Markdown lessons. "
          "Full lesson text is included in the notes of each lesson divider; selected SQL/DAX appears "
          "on teaching slides. Live demonstrations require target-tenant preflight. The earlier lesson "
          "plan flags a possible 1:30-2:00 PM conflict: resolve it or arrange a handoff before delivery.")

    table("Our three-hour learning path", ["Lesson", "Focus", "Minutes"], [
        ["01", "Architecture decision framework", "25"],
        ["02", "Repeatable ingestion", "35"],
        ["03", "Curated SQL engineering", "50"],
        ["04", "Consumption architecture", "35"],
        ["05", "Operations and GCC High findings", "25"],
        ["06", "Exit review", "10"],
    ], [1.15, 9.2, 1.65],
          speaker="180 minutes total. Introductions and transitions fit inside these module allocations. "
          "The SQL objects should be prepared before class; do not type the entire procedure live.")

    cards("The standard: explain it, repeat it, support it", [
        ("Explain it", "Defend the architecture.\nPreserve source evidence.\nExplain every exception."),
        ("Repeat it", "Reload without duplicates.\nGate dependent steps.\nProve the intended publication."),
        ("Support it", "Name owners.\nTest consumer access.\nSeparate skills from environment issues."),
    ], "General product documentation is not proof of GCC High tenant availability.",
          speaker="The session is SQL/Power Query-first. Notebooks are discussed only when a requirement "
          "justifies them. Prior successful ingestion and prior pipeline authentication failures are "
          "historical observations, not universal capability statements.")

    section(1, "Architecture decision framework", 25, "A defensible architecture decision record")
    cards("Start with the business question", [
        ("Compare", "Do two extracts describe the same records, amounts, and statuses?"),
        ("Investigate", "Which differences or data-quality problems require attention?"),
        ("Repeat", "Can the next comparison run without hidden manual repairs?"),
    ], "Moving files is an activity. Reconciliation is the outcome.",
          speaker="Revisit Session 2. Both files come from SharePoint. Column names and value "
          "representations differ. The audience needs both a summary and investigable exceptions.")
    cards("Define what one result row means", [
        ("Grain", "One normalized business record ID in one comparison of two complete snapshots."),
        ("Implication", "Duplicate source rows are ambiguity to report, not extra business records."),
    ], "The output grain determines the join and the quality rules.",
          speaker="Ask why eleven input rows need not produce eleven output rows. Our baseline has "
          "six distinct keys. Do not select a random duplicate to force uniqueness.")
    table("Warehouse or Lakehouse?", ["Decision", "Warehouse", "Lakehouse"], [
        ["Development", "Primarily T-SQL", "Often Spark / file engineering"],
        ["Data focus", "Relational analytical tables", "Files, Delta, mixed data types"],
        ["T-SQL data writes", "Supported Warehouse DML", "Not through SQL analytics endpoint"],
        ["SQL transactions", "Multi-table transaction support", "Do not assume the same model"],
        ["Team fit", "SQL-oriented engineering", "Appropriate Lakehouse write skills"],
    ], [2.35, 4.55, 5.1],
          speaker="Both use OneLake and Delta. Lakehouse SQL query access is not a Warehouse write path. "
          "Neither choice implies all SQL Server features are supported.", size=17,
          takeaway_text="Choose the development and operating model, not a product popularity contest.")
    bullets("A Warehouse-centered path for this workload", [
        "SharePoint extracts -> Dataflow Gen2",
        "Warehouse landing tables -> SQL normalization and quality checks",
        "SQL reconciliation -> curated Gold table",
        "Shared semantic model -> governed Power BI consumption",
    ], "Add another engineering component only when a requirement justifies it.",
          speaker="Tabular inputs, relational comparison, existing SQL/Power Query skills. "
          "A notebook is not a maturity badge. New custom Python extraction could justify revisiting "
          "part of the design without replacing everything.")
    cards("Bronze, Silver, Gold: responsibilities first", [
        ("BRONZE", "Preserve source-aligned values, duplicates, and invalid input evidence."),
        ("SILVER", "Align names and representations.\nEvaluate key and value quality."),
        ("GOLD", "Publish the business result and investigation fields for approved consumers."),
    ], "Three responsibilities do not automatically require three workspaces.",
          speaker="The training implementation uses schemas in one Warehouse and views where appropriate. "
          "Production isolation follows security, lifecycle, performance, and ownership requirements.")
    exercise("Decision exercise", "Choose a starting point. Then introduce a new requirement: "
             "documents need custom Python extraction.",
             "A decision record: outcome, grain, choice, rationale, alternative, assumptions, revisit condition.",
             "Warehouse is justified for the baseline. Identify precisely which part of the new requirement "
             "needs a different tool. Do not automatically redesign every stage.",
             speaker="Use the worked architecture decision record in Lesson 1 as a reference. "
             "Ask participants to name the rejected alternative and why it is not needed now.")

    section(2, "Repeatable ingestion", 35, "An ingestion contract and a documented refresh method")
    cards("A folder is not a source contract", [
        ("Scope", "Exact site, folder, files, columns, and comparison period."),
        ("Readiness", "How the source owner confirms a complete matching snapshot pair."),
        ("Execution", "Approved connection, intended execution identity, and known destination behavior."),
    ], "File presence alone does not prove that two extracts belong together.",
          speaker="Exclude unrelated and temporary files. A browser session under the presenter does "
          "not prove scheduled connection access. Validate gateway dependencies only where applicable.")
    table("System A: preserve the intentional problems", ["RecordId", "Amount", "Status"], [
        ["A001", "100.00", "Active"], ["A002", "200.00", "Active"],
        ["A003", "300.00", "Active"], ["A005", "500.00", "Active"],
        ["A005", "500.00", "Active"], ["A006", "not provided", "Active"],
    ], [3.2, 5.6, 3.2],
          takeaway_text="Six rows. The duplicate and invalid amount must survive landing.",
          speaker="Synthetic system_a.csv. All source columns land as text. A005 occurs twice and "
          "A006 has a deliberately unconvertible amount.")
    table("System B: representations and business differences", ["Reference", "Total", "State"], [
        ['" a001 "', "100.00", "ACTIVE"], ["A002", "250.00", "ACTIVE"],
        ["A004", "400.00", "ACTIVE"], ["A005", "500.00", "ACTIVE"],
        ["A006", "600.00", "ACTIVE"],
    ], [3.2, 5.6, 3.2],
          takeaway_text="Five rows. Quotes above expose A001's surrounding spaces; they are not CSV data.",
          speaker="Synthetic system_b.csv. The actual first Reference value contains spaces and lowercase "
          "characters. A002 differs in amount; A003 and A004 occur on only one side.")
    bullets("Dataflow walkthrough: select, preserve, land", [
        "Connect through the approved SharePoint connection and execution identity.",
        "Filter the exact folder and file; inspect expected columns.",
        "Inspect automatic type changes; keep landing values as text.",
        "Map each query to its explicit existing Warehouse landing table.",
        "Run, inspect refresh history, and verify destination counts: 6 and 5.",
    ], "A Power Query preview is not evidence of a completed destination load.",
          speaker="Prepare bronze.system_a and bronze.system_b before class using Lesson 3 setup. "
          "Do not remove invalid rows to make refresh green. Inspect the actual destination mapping "
          "and supported settings in the target environment.")
    cards("Replace, append, and incremental are different", [
        ("REPLACE", "The incoming complete snapshot becomes the current state.\nRerun: 6 and 5 rows."),
        ("APPEND", "Incoming rows are added.\nSame snapshots twice: 12 and 10 rows."),
        ("INCREMENTAL", "Changes need keys, reliable change evidence, updates/deletes, checkpoints, and retry rules."),
    ], "An append switch is not an incremental correctness strategy.",
          speaker="Explain updates, source deletions, timestamp ties, and late-arriving changes. "
          "Full replacement is appropriate for this small teaching example, not every enterprise source.")
    cards("Scheduling starts work. Orchestration gates it.", [
        ("Schedule", "Start ingestion at a configured time.\nA time gap does not prove success."),
        ("Orchestration", "Wait for both intended loads, then validate, reconcile, and update model freshness."),
    ], "System A today + System B yesterday is not a valid comparison.",
          speaker="Independent destination replacements are not an atomic snapshot-pair update. "
          "The lab can use an explicit supervised gate. If pipeline invocation is blocked, direct "
          "ingestion scheduling plus manual SQL remains partial automation.")
    table("Give each tool a defined job", ["Tool", "Responsibility in this design"], [
        ["Dataflow Gen2", "Power Query selection, shaping, and loading"],
        ["Pipeline", "Supported activity coordination and success/failure dependencies"],
        ["Copy job", "Consider for supported movement-oriented requirements"],
        ["Warehouse SQL", "Reusable relational quality and reconciliation rules"],
        ["Notebook", "Use when a justified pro-code requirement exists"],
    ], [3, 9],
          speaker="Validate each connector, destination, load mode, and Government capability. "
          "A copy job is not assumed to replace this SharePoint/Power Query transformation.")
    exercise("Prove the refresh contract", "Load the unchanged snapshots twice. What must remain "
             "the same, and what stops if System B fails?",
             "Source scope, mapping, replace behavior, counts, trigger, execution owner, and downstream gate.",
             "Landing remains 6 and 5, including the original duplicate. A failed source blocks "
             "new reconciliation. Restore a complete matching pair before resuming.")

    section(3, "Curated SQL engineering", 50, "Curated Warehouse objects and explainable reconciliation results")
    cards("Build a chain of explainable SQL objects", [
        ("BRONZE TABLES", "system_a / system_b\nSource-aligned text columns and original row evidence."),
        ("SILVER VIEWS", "normalized_records\nrecord_summary\nShared normalization and quality evidence."),
        ("GOLD TABLE", "reconciliation\nOne current result per key, published by a stored procedure."),
    ], "Run only in the isolated training Warehouse, not a Lakehouse SQL analytics endpoint.",
          speaker="Create schemas once and execute object definitions separately. Preflight examples. "
          "The full SQL is in Lesson 3 and in the notes on this module divider. "
          "Prepare tables before the Lesson 2 ingestion walkthrough.")
    code("Normalize values without losing evidence",
         "NULLIF(UPPER(TRIM(RecordId)), '')\n"
         "    AS record_id,\n\n"
         "TRY_CONVERT(\n"
         "    decimal(18,2),\n"
         "    NULLIF(TRIM(Amount), '')\n"
         ") AS amount",
         "A common vocabulary",
         "Map both source schemas.\nKeep raw values.\nBlank keys become null.\nInvalid amounts do not become zero.",
         speaker="This is an excerpt from silver.normalized_records, not a complete runnable statement. "
         "System B uses Reference/Total/State. Uppercasing is a domain-approved teaching rule, "
         "not a universal identifier rule. Define currency, precision, locale, and tolerance for real data.")
    cards("A view is logic, not a frozen snapshot", [
        ("Reuse", "Silver views expose common names and consistent normalization without another loaded copy."),
        ("Control", "Views reflect underlying Bronze changes. Serialize source loading and reconciliation."),
    ], "A usable decimal value and a trustworthy source comparison are separate conditions.",
          speaker="Query raw and normalized fields side by side. A006's raw text explains its null amount. "
          "Converting to decimal(18,2) does not prove the source originally had exactly two decimals.")
    code("Check key quality before joining",
         "GROUP BY source_system, record_id\n\n"
         "COUNT(*) AS source_row_count\n\n"
         "CASE WHEN COUNT(*) = 1\n"
         "     THEN MAX(amount)\n"
         "END AS amount",
         "Do not invent a winner",
         "Summarize each source/key.\nExpose a value only for one row.\nKeep duplicate counts.\nBlock blank-key publication.",
         speaker="Illustrative clauses from silver.record_summary, shown out of full SELECT context. "
         "MAX is guarded by COUNT(*) = 1; it is not used to choose among duplicates. "
         "The full view includes invalid_amount_count and invalid_status_count. Blank keys remain "
         "visible in normalized_records and cause the teaching procedure to fail.")
    table("Classification precedence is a business rule", ["Order", "First applicable condition", "Result"], [
        ["1", "Duplicate key in either source", "DUPLICATE_KEY"],
        ["2", "Unusable amount / status", "INVALID_AMOUNT / INVALID_STATUS"],
        ["3", "Key absent from one side", "ONLY_IN_A / ONLY_IN_B"],
        ["4", "Usable values differ", "Amount / status / combined mismatch"],
        ["5", "Both usable records agree", "MATCH"],
    ], [1.15, 5.25, 5.6], size=17,
          takeaway_text="A primary classification does not erase other source-quality evidence.",
          speaker="Amount-invalid precedes status-invalid. Duplicates take precedence even when values "
          "agree. Existing invalid data precedes missing-counterpart classifications. All rules are "
          "explicit in the procedure and require business approval for production.")
    cards("An outer join preserves the business question", [
        ("INNER JOIN", "Only keys found on both sides survive.\nSource-only exceptions disappear."),
        ("FULL OUTER JOIN", "Keys from either side remain visible.\nCompare summarized source/key rows."),
    ], "Join after quality assessment so duplicate source rows do not multiply the result.",
          speaker="A003 exists only in A; A004 only in B. Both must be reported. "
          "A005 produces one DUPLICATE_KEY output, not two supposed business records.")
    code("A procedure defines repeatable SQL work",
         "EXEC gold.usp_refresh_reconciliation\n"
         "    @BatchId =\n"
         "      'training-2026-10-28-baseline';",
         "Definition vs. execution",
         "CREATE PROCEDURE defines the steps.\nEXEC performs them.\nThe batch label is not proof that sources align.",
         speaker="Walk through the complete procedure in the lesson. It validates nonempty inputs and "
         "blank keys, replaces Gold, checks output uniqueness, and rethrows failures. "
         "Confirm completed matching sources and no competing refresh before execution.")
    bullets("Protect publication, not just execution", [
        "Begin a transaction and validate the teaching inputs.",
        "Replace the current Gold result using the explicit rules.",
        "Check one output row per key before committing.",
        "On error: roll back the replacement and rethrow to the caller.",
    ], "An atomic Gold replacement does not make upstream dataflows atomic.",
          speaker="The example rejects empty teaching inputs and blank keys, while publishing duplicates "
          "and invalid values as exceptions. These policies are not universal. It overwrites current "
          "Gold rather than retaining history. Production needs durable source/batch controls and "
          "concurrency policy. Do not run overlapping refreshes.")
    table("Expected baseline: every result is explainable", ["Key", "Classification", "Why"], [
        ["A001", "MATCH", "Formatting normalizes; values agree"],
        ["A002", "AMOUNT_MISMATCH", "200.00 vs. 250.00"],
        ["A003", "ONLY_IN_A", "No corresponding B record"],
        ["A004", "ONLY_IN_B", "No corresponding A record"],
        ["A005", "DUPLICATE_KEY", "Two A rows; no chosen winner"],
        ["A006", "INVALID_AMOUNT", '"not provided" is not a decimal'],
    ], [1.1, 4.65, 6.25], size=17,
          takeaway_text="11 input rows -> 6 business keys -> 1 match + 5 exceptions.",
          speaker="Ask learners to explain a row using source evidence. Do not interpret five exceptions "
          "as a failed lab; they are the intended synthetic result.")
    exercise("Rerun, then follow one correction", "Run unchanged input twice. Then change B/A002 "
             "from 250.00 to 200.00 and predict the result.",
             "Observed rerun counts and a source-to-Gold explanation of the correction.",
             "Unchanged input remains six rows. The correction produces two matches and four exceptions. "
             "Timestamps can change. Reload source files rather than patching Gold. Restore baseline "
             "before later demonstrations that use the original totals.")

    section(4, "Consumption architecture", 35, "A governed consumption path with a documented rationale")
    cards("Separate shared semantics from storage mode", [
        ("Shared semantic model", "Reusable measures, names, relationships, formatting, and model access rules."),
        ("Storage / query mode", "How model tables obtain and query data: Import, DirectQuery, or Direct Lake."),
    ], "A shared semantic model can use Direct Lake; they are not competing choices.",
          speaker="Record-level reconciliation classification remains in Gold for every consumer. "
          "Reporting measures can live in the shared model. Do not rebuild the algorithm per report.")
    table("Choose a verified consumption path", ["Route", "Behavior", "Verify"], [
        ["Import", "Copies data into model storage", "Successful refresh after publication"],
        ["DirectQuery", "Queries the source through the model", "Performance, identity, connection"],
        ["Direct Lake", "Uses supported OneLake Delta tables", "Objects, framing, identity, capacity"],
        ["Direct SQL", "SQL client reads curated Warehouse", "SQL permissions and shared definitions"],
    ], [2.0, 5.0, 5.0], size=17,
          speaker="Direct SQL is a separate route, not a semantic-model storage mode. "
          "Direct Lake uses memory; it is not 'no refresh or freshness management.' "
          "Import can be the correct verified choice.")
    cards("Direct Lake: verify the actual variant", [
        ("On SQL", "SQL endpoint discovery and permission behavior.\nCan fall back to DirectQuery under supported settings."),
        ("On OneLake", "Uses compatible Delta sources.\nDoes not use the same DirectQuery fallback path."),
    ], "Use the physical Gold table; do not assume SQL views behave like Direct Lake tables.",
          speaker="A SQL view can trigger fallback for Direct Lake on SQL, depending on behavior settings. "
          "Do not assume either variant is available in the target Government tenant. Inspect the model "
          "actually created and its source path.")
    code("Create measures that reuse the business contract",
         "Compared Records =\n"
         "COUNTROWS('Reconciliation')\n\n"
         "Match Rate =\n"
         "DIVIDE(\n"
         "    [Matched Records],\n"
         "    [Compared Records]\n"
         ")",
         "Count business keys",
         "Name the model table Reconciliation.\nDefine exception as not MATCH.\nKeep denominator and filters explicit.",
         speaker="Map the table to gold.reconciliation. Format Match Rate as a percentage. "
         "The model counts published keys, not eleven source rows. Gold contains only the latest "
         "publication; its timestamp does not create historical trend data.")
    code("Respect the report's filter context",
         "Exception Records =\n"
         "COALESCE(\n"
         "    CALCULATE(\n"
         "        [Compared Records],\n"
         "        KEEPFILTERS(\n"
         "          'Reconciliation'[classification]\n"
         "            <> \"MATCH\"\n"
         "        )\n"
         "    ),\n"
         "    0\n"
         ")",
         "One definition",
         "KEEPFILTERS intersects existing selections.\nMatched Records uses the same pattern with = MATCH.",
         speaker="Use the full measures from Lesson 4. A separate metric that excludes incomparable "
         "records requires a different approved denominator and name. Do not hide quality exceptions "
         "to improve the reported match rate.")
    slide = base("A report should summarize and explain")
    for i, (number, label) in enumerate([
        ("6", "Compared"), ("1", "Matched"), ("5", "Exceptions"), ("16.67%", "Match rate")
    ]):
        x = 0.65 + i * 3.06
        rect(slide, x, 2.0, 2.82, 1.7, WHITE)
        text(slide, number, x + 0.15, 2.19, 2.52, 0.68, 37, TEAL, True)
        text(slide, label, x + 0.15, 3.02, 2.52, 0.4, 16, GRAY)
    text(slide, "From a summary to an explanation", 0.75, 4.15, 11.7, 0.55, 24, INK, True)
    text(slide, "A002  |  A amount 200.00  |  B amount 250.00  |  AMOUNT_MISMATCH\n"
         "Show classification detail, source values, batch ID, and publication time.",
         0.75, 4.95, 11.7, 1, 20)
    takeaway(slide, "Illustrative baseline, not a screenshot or evidence of a live tenant execution.")
    notes(slide, "Demonstrate the prepared report using the actual model mode. Trace A002 to SQL. "
          "Trace A005 as a quality exception. These cards are illustrative and editable PowerPoint "
          "shapes, not claimed screenshots of a deployed report.")
    cards("Govern each access path independently", [
        ("Report / model", "View vs. Build permissions.\nShared measures and model rules.\nTest the intended reader."),
        ("SQL / OneLake", "Separate object and route permissions.\nDAX and model RLS do not automatically protect other paths."),
    ], "An administrator's successful view is not a consumer-access test.",
          speaker="Direct SQL readers should use curated Gold rather than rebuilding from Bronze. "
          "If an aggregate definition must be shared in SQL, provide an approved SQL contract. "
          "Avoid broad workspace write access for readers.")
    exercise("Choose and prove the consumption path", "Which audience uses the shared model, which "
             "needs SQL, and how do you prove both see the intended publication?",
             "Primary route, actual mode, definition owners, approved access, and freshness evidence.",
             "Verify source load, SQL publication, model refresh/framing/query path, and consumer result. "
             "Use batch ID, publication time, and known totals together. Record the model/access owner.")

    section(5, "Operations and GCC High findings", 25, "An operations checklist and separate findings register")
    bullets("Use success gates, not optimistic time gaps", [
        "Confirm a complete, matching source pair.",
        "Confirm ingestion and quality/batch checks.",
        "Complete SQL publication, then the mode-specific model freshness step.",
        "Verify the consumer result and record the run outcome.",
    ], "Scheduled ingestion + supervised downstream steps is partial automation.",
          speaker="If pipeline invocation remains blocked, state the operating dependency honestly. "
          "A 6:00 ingestion and 6:15 SQL clock do not guarantee sequencing. "
          "No live path is proven by a diagram alone.")
    table("Monitor execution and business correctness", ["Evidence", "What it proves", "Still check"], [
        ["Activity status", "Configured work completed", "Source scope and business rules"],
        ["Counts / quality", "Volume and key conditions", "Batch alignment and access"],
        ["Publication / consumer", "Intended output is visible", "Owners and other control gates"],
    ], [2.8, 4.5, 4.7],
          takeaway_text="A green activity is necessary evidence, not the whole definition of success.",
          speaker="Record run ID, source evidence, times, object names, actual/expected counts, result, "
          "and last successful publication. Synthetic counts are not universal production thresholds. "
          "Sanitize errors and never record credentials or tokens.")
    table("Recover at the right boundary", ["Failure", "Response"], [
        ["System B ingestion", "Stop new comparison; restore and verify a matching pair"],
        ["SQL before commit", "Roll back Gold replacement; investigate SQL/data evidence"],
        ["Model freshness", "Distinguish new Gold from stale model; repair the model stage"],
        ["Unexpected report total", "Inspect filters, rules, batch, and publication before escalation"],
    ], [3.1, 8.9],
          speaker="The prior report remaining accessible does not mean today's run succeeded. "
          "Do not rerun every stage automatically. A retry needs known idempotency and concurrency rules.")
    cards("Access is not the same as ownership", [
        ("Workspace roles", "Admin / Member / Contributor / Viewer govern collaboration.\nDo not elevate readers to fix access."),
        ("Artifact permissions", "Warehouse, model Build, report access, and connections have distinct responsibilities."),
        ("Named owners", "Source, connection, ingestion, SQL/data, model/report, support.\nInclude backup coverage."),
    ], "Test restricted consumers without broad workspace editing privileges.",
          speaker="Exact role capabilities depend on current item behavior. Elevated workspace "
          "editing users are not suitable tests of restricted model RLS. One person may hold "
          "multiple owner roles, but each duty needs explicit coverage.")
    bullets("Naming and lifecycle make change supportable", [
        "Use approved naming conventions; teaching schemas are examples.",
        "Review rule changes against known results and failure behavior.",
        "Promote through the approved, verified mechanism; retain recovery options.",
        "Check shared dependencies and owner approval before cleaning up assets.",
    ], "A name containing 'test' does not provide production isolation.",
          speaker="Changing invalid amounts from published exceptions to blockers affects SQL, alerts, "
          "report freshness, and the runbook. Keep them consistent. Do not assume commercial Git "
          "or deployment features are available in Government.")
    cards("Route the issue before proposing the fix", [
        ("LEARNING NEED", "Cannot explain duplicate handling or procedure execution.\nAssign focused practice."),
        ("DATA ISSUE", "Duplicate key, invalid amount, missing extract.\nRoute to the data-quality/source owner."),
        ("ENVIRONMENT", "Reproducible connection, permission, gateway, service, or capability behavior.\nInvestigate scope and cause."),
    ], "An authentication error does not establish a GCC High-wide unsupported feature.",
          speaker="Keep learning and environment logs separate; route data exceptions through the "
          "data-quality process. Findings need scope, time, role, repro, expected/actual behavior, "
          "sanitized evidence, impact, investigation state, owner, and next action. "
          "Document interim approach, manual effort, dependency gaps, and review condition.")
    exercise("System A succeeded. System B failed.", "Yesterday's report still opens. Has today's "
             "comparison succeeded? What do you do next?",
             "An ordered operating checklist plus an owned, evidence-based finding.",
             "No. Stop new publication, preserve and identify the last successful result, record the "
             "failed stage, and assign investigation. Resume only after a verified complete matching "
             "source pair. Include SQL/model/access/freshness gates and owners in the checklist.")

    section(6, "Exit review", 10, "An honest completion record, owned actions, and follow-up track choices")
    table("Inspect the outputs, not just attendance", ["Artifact", "Acceptance evidence"], [
        ["Architecture decision", "Rationale, assumptions, alternative, revisit condition"],
        ["Ingestion contract", "Source scope, mode, counts, trigger, downstream gate"],
        ["SQL result", "Explained classifications and observed or labeled rerun"],
        ["Consumption decision", "Actual mode, permissions, freshness"],
        ["Operations / findings", "Success gates, failure handling, and named owners"],
    ], [3.5, 8.5],
          speaker="Use the first four minutes for artifact inspection and teach-back. "
          "Ask learners to explain A005, A006, reruns, and what a dataflow schedule does not prove.")
    cards("Describe the evidence honestly", [
        ("Completed / demonstrated", "Participant completed: performed and explained.\nPresenter demonstrated: shown live, not independently completed."),
        ("Reviewed / blocked", "Prepared evidence reviewed: not a live proof.\nBlocked: a required dependency remains unresolved."),
    ], "These are evidence states, not grades.",
          speaker="Prepared evidence can preserve teaching continuity but cannot prove the participant's "
          "connection works. Ask whether another practitioner could repeat the scenario without hidden steps.")
    bullets("Turn open questions into owned work", [
        "Name the specific unresolved behavior or skill.",
        "Agree an owner and next action; do not invent commitments.",
        "Define completion evidence and an agreed review point.",
        "Keep expected data exceptions separate from implementation defects.",
    ], "A question without an owner is not an action plan.",
          speaker="Use approximately four minutes for follow-ups. Replace 'pipelines are broken' with "
          "a scoped retest using the intended connection and role, a sanitized run result, and a finding "
          "update. Use item/state/owner/action/evidence/review-point fields.")
    cards("Choose follow-up tracks by responsibility", [
        ("SQL / WAREHOUSE", "Relational modeling, procedures, full/incremental loads, data quality, SQL operations."),
        ("NOTEBOOK / SPARK", "Python/PySpark, Delta engineering, parameterized notebooks, pro-code orchestration."),
    ], "Confirm each participant's track and next skill; neither track is a maturity ranking.",
          speaker="Use the final two minutes. Participants can have a primary track and later learn "
          "the other. Confirm choices with participants and the technical lead. Today's session is "
          "a learning implementation, not a claim that all production-readiness gates have passed.")
    slide = base("One workload. A repeatable delivery model.", dark=True)
    text(slide, "Decide. Ingest. Curate. Consume. Operate.",
         0.7, 2.55, 11.9, 1.2, 35, WHITE, True)
    text(slide, "The completion test:\nAnother practitioner can repeat the scenario using the "
         "agreed artifacts, without hidden steps.",
         0.7, 4.2, 11.9, 1.65, 27, MINT)
    notes(slide, "Close with the artifacts and owned work. The six source lessons remain the detailed "
          "reference. Do not imply unattended automation, consumer permissions, or Government "
          "availability has been validated unless it was actually observed.")

    module = "REFERENCE"
    lesson = None
    slide = base("Sources and delivery boundaries")
    references = [
        ("Warehouse / Lakehouse decision guide",
         "https://learn.microsoft.com/en-us/fabric/fundamentals/decision-guide-lakehouse-warehouse"),
        ("Dataflow Gen2 destinations",
         "https://learn.microsoft.com/en-us/fabric/data-factory/dataflow-gen2-data-destinations-and-managed-settings"),
        ("Warehouse T-SQL surface area",
         "https://learn.microsoft.com/en-us/fabric/data-warehouse/tsql-surface-area"),
        ("Warehouse transactions",
         "https://learn.microsoft.com/en-us/fabric/data-warehouse/transactions"),
        ("Direct Lake overview",
         "https://learn.microsoft.com/en-us/fabric/fundamentals/direct-lake-overview"),
    ]
    for i, (label, url) in enumerate(references):
        box = text(slide, label, 0.75, 1.95 + i * 0.72, 11.8, 0.5, 21, TEAL)
        box.text_frame.paragraphs[0].runs[0].hyperlink.address = url
    takeaway(slide, "Validate live paths in the target GCC High tenant; documentation is not tenant evidence.")
    notes(slide, "Deck sources:\n" + "\n".join(path.name for path in SOURCES)
          + "\n\nProduct references:\n" + "\n".join(url for _, url in references)
          + "\n\nFull Markdown lesson content is embedded in the notes of the six lesson divider slides. "
          "Code on individual teaching slides may be an excerpt; use the complete lesson for execution. "
          "Live demonstrations, schedules, credentials, permissions, and model modes require preflight.")


if __name__ == "__main__":
    if len(SOURCES) != 6:
        raise SystemExit(f"Expected six source lessons, found {len(SOURCES)}")
    build()
    OUTPUT.parent.mkdir(exist_ok=True)
    prs.save(OUTPUT)
    print(f"Created {OUTPUT}")
    print(f"{len(prs.slides)} slides; all have speaker notes.")
    print(f"Smallest non-footer text: "
          f"{min(size for _, size, _ in layout_records if size >= 14)} pt")
